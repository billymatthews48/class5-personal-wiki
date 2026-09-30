"""Read raw source files into passages. Raw files are opened read-only and never modified.

Each passage keeps a locator (section heading, page or slide) so an answer can cite the exact
place in the original file. Everything here is local parsing; no model is involved.
"""
import re
from dataclasses import dataclass, field
from pathlib import Path

from .. import config


@dataclass
class Segment:
    locator: str   # "section 'Lean Operations'", "page 4", "slide 12"
    text: str


@dataclass
class Document:
    raw_path: str              # vault-relative, forward slashes: "raw/mba/Marketing/Hwk 3.docx"
    kind: str                  # docx | pdf | pptx | markdown
    title_hint: str            # file stem; used only as a hint, never as a note name
    segments: list = field(default_factory=list)

    @property
    def text(self):
        return "\n\n".join(s.text for s in self.segments)


READERS = {}


def reader(*suffixes):
    def wrap(fn):
        for s in suffixes:
            READERS[s] = fn
        return fn
    return wrap


def supported(path):
    p = Path(path)
    return p.suffix.lower() in READERS and not p.name.startswith("~$")


def read(path):
    """Parse one raw file into a Document (or None if it has no extractable text)."""
    path = Path(path).resolve()
    items = READERS[path.suffix.lower()](path)
    if not items:
        return None
    rel = path.relative_to(config.VAULT.resolve()).as_posix() if path.is_relative_to(config.VAULT.resolve()) else str(path)
    return Document(raw_path=rel, kind=path.suffix.lower().lstrip("."),
                    title_hint=path.stem, segments=pack(items))


def discover(path):
    path = Path(path)
    files = [path] if path.is_file() else sorted(p for p in path.rglob("*") if p.is_file())
    for f in files:
        if supported(f):
            doc = read(f)
            if doc:
                yield doc


# ---------- readers: each returns a list of (locator, text) ----------

@reader(".docx")
def _docx(path):
    import docx
    d = docx.Document(str(path))
    items, heading = [], None
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading") or style == "title":
            heading = text
        items.append((f"section '{heading}'" if heading else "opening", text))
    for ti, table in enumerate(d.tables, 1):
        for row in table.rows:
            cells = []
            for c in row.cells:   # merged cells repeat; keep each value once
                t = c.text.strip()
                if t and t not in cells:
                    cells.append(t)
            if cells:
                items.append((f"table {ti}", " | ".join(cells)))
    return items


@reader(".pdf")
def _pdf(path):
    from pypdf import PdfReader
    items = []
    pages = PdfReader(str(path)).pages
    for i, page in enumerate(pages, 1):
        # PDF line breaks are layout, not meaning (some exports put one word per line), so
        # a page becomes one flowing block that pack() re-splits on sentence boundaries.
        text = re.sub(r"\s+", " ", page.extract_text() or "").strip()
        if text:
            items.append((f"page {i}", text))
    if not items:            # scanned PDF: no text layer at all -> local OCR
        items = _ocr_pdf(path)
    return items


TESSERACT = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")


def _ocr_pdf(path):
    """Render each page with pypdfium2 and read it with local Tesseract. Offline; no cloud OCR."""
    try:
        import pypdfium2 as pdfium
        import pytesseract
    except ImportError:
        return []
    if TESSERACT.exists():
        pytesseract.pytesseract.tesseract_cmd = str(TESSERACT)
    items = []
    pdf = pdfium.PdfDocument(str(path))
    for i in range(len(pdf)):
        image = pdf[i].render(scale=2).to_pil()
        data = pytesseract.image_to_data(image, output_type=pytesseract.Output.DICT)
        words = [(w, float(c)) for w, c in zip(data["text"], data["conf"]) if w.strip() and float(c) >= 0]
        conf = sum(c for _, c in words) / len(words) if words else 0
        text = re.sub(r"\s+", " ", " ".join(w for w, _ in words)).strip()
        OCR_LOG.append({"file": path.name, "page": i + 1, "mean_conf": round(conf, 1),
                        "kept": conf >= OCR_MIN_CONF and len(text) > 40})
        # Handwritten scans OCR into noise (low confidence); indexing them would pollute search.
        if conf >= OCR_MIN_CONF and len(text) > 40:
            items.append((f"page {i + 1} (OCR)", text))
    return items


OCR_MIN_CONF = 70
OCR_LOG = []   # per-page OCR confidence, reported by ingest


@reader(".pptx", ".pptm")
def _pptx(path):
    from pptx import Presentation
    items = []
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        texts = [sh.text_frame.text.strip() for sh in slide.shapes
                 if sh.has_text_frame and sh.text_frame.text.strip()]
        if texts:
            items.append((f"slide {i}", "\n".join(texts)))
    return items


@reader(".md", ".txt")
def _markdown(path):
    items, heading = [], None
    for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        if not line.strip():
            continue
        if line.startswith("#"):
            heading = line.lstrip("#").strip()
        items.append((f"section '{heading}'" if heading else "opening", line.rstrip()))
    return items


# ---------- passage packing ----------

def pack(items, size=config.CHUNK_CHARS, overlap=config.CHUNK_OVERLAP):
    """Pack consecutive (locator, text) items into ~size-char passages.

    A passage never spans two locators with different pages/slides unless they are tiny, so
    citations stay precise. Over-long items are split on sentence boundaries with overlap.
    """
    segs, buf, loc = [], [], None
    def flush():
        if buf:
            segs.append(Segment(loc, "\n".join(buf)))
    for locator, text in items:
        for piece in _split(text, size, overlap):
            if buf and (locator != loc or sum(map(len, buf)) + len(piece) > size):
                if locator != loc and sum(map(len, buf)) < size // 4 and len(piece) < size // 2:
                    if locator not in loc.split("; "):          # merge a tiny leftover
                        loc = f"{loc}; {locator}"
                else:
                    flush()
                    buf, loc = [], None
            if loc is None:
                loc = locator
            buf.append(piece)
    flush()
    return segs


def _split(text, size, overlap):
    if len(text) <= size:
        return [text]
    sentences = re.split(r"(?<=[.!?])\s+", text)
    out, cur = [], ""
    for s in sentences:
        while len(s) > size:           # a single enormous "sentence" (tables, bullet dumps)
            out.append(s[:size])
            s = s[size - overlap:]
        if cur and len(cur) + len(s) + 1 > size:
            out.append(cur)
            cur = cur[-overlap:] + " " + s
        else:
            cur = f"{cur} {s}".strip()
    if cur:
        out.append(cur)
    return out
